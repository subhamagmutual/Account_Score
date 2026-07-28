"""Generate PAS_Code_Guide.xlsx -- a reference map of the codebase."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, sz=10, bold=True, color="FFFFFF")
BODY = Font(name=FONT, sz=10)
BODY_B = Font(name=FONT, sz=10, bold=True)
TITLE = Font(name=FONT, sz=14, bold=True, color="1F3864")
SUB = Font(name=FONT, sz=10, italic=True, color="595959")
MONO = Font(name="Consolas", sz=9)
ENTRY_FILL = PatternFill("solid", fgColor="FFF2CC")   # highlights entry points
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()


def sheet(name, title, subtitle=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws.row_dimensions[1].height = 22
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = SUB
        ws.row_dimensions[2].height = 15
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, cols):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, start_row, rows, widths, wrap_cols=(), fills=None):
    for r, vals in enumerate(rows, start=start_row):
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY
            cell.border = BOX
            cell.alignment = Alignment(
                vertical="top", wrap_text=(c in wrap_cols))
            if fills:
                f = fills(r - start_row, c, vals)
                if f:
                    cell.fill = f
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_table(ws, name, first_row, last_row, last_col):
    ref = f"A{first_row}:{get_column_letter(last_col)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)


# ===========================================================================
# 1. START HERE
# ===========================================================================
ws = sheet("Start Here", "Physician MPL PAS — code guide",
           "Which file to run, in what order, and what to expect. "
           "Entry points are highlighted in amber.")

ws["A4"] = "THE ENTRY POINT"
ws["A4"].font = Font(name=FONT, sz=11, bold=True)
ws["A5"] = ("PAS_Pipeline.ipynb is the entry point for analysing your PAS model. "
            "It is the only notebook in the project.")
ws["A5"].font = BODY_B
ws["A6"] = ("The other two runnable files are tests, not analysis. Run them once to "
            "confirm the harness works, then work in the notebook.")
ws["A6"].font = BODY
for r in (5, 6):
    ws.merge_cells(f"A{r}:F{r}")
    ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[5].height = 26
ws.row_dimensions[6].height = 26

header(ws, 8, ["Step", "Run this", "Command", "What it proves",
               "Expected result", "How often"])
steps = [
    (1, "test_edges.py", "python test_edges.py",
     "The harness code itself is correct — 32 adversarial tests covering empty "
     "frames, all-NaN columns, duplicate indices, constant targets, both CLI "
     "exit codes.",
     "32 passed, 0 failed", "Once after install, and after any code change"),
    (2, "make_fixture.py",
     "python make_fixture.py",
     "Builds 60,000 synthetic rows with six bugs planted from your design "
     "workbook, so you can confirm the checks actually fire.",
     "Writes fixture_physician_scores.csv", "Once, then whenever checks change"),
    (3, "pas_diagnostics (CLI)",
     "python -m pas_diagnostics --input fixture_physician_scores.csv --outdir out",
     "The checks detect the planted bugs. If this returns zero blockers, "
     "nothing downstream is trustworthy.",
     "9 BLOCKER / 27 HIGH / 16 MEDIUM", "Once, as a smoke test"),
    (4, "PAS_Pipeline.ipynb",
     "Open in Jupyter / VS Code, run cells top to bottom",
     "THIS IS THE ACTUAL WORK. Points verified code at your real scored output "
     "to find out whether the PAS score sorts loss.",
     "Depends on your data — that is the point",
     "Ongoing. This is where you live."),
]
body(ws, 9, steps, [6, 22, 52, 52, 30, 34], wrap_cols=(2, 3, 4, 5, 6),
     fills=lambda i, c, v: ENTRY_FILL if i == 3 else None)
for r in range(9, 13):
    ws.row_dimensions[r].height = 62
add_table(ws, "RunOrder", 8, 12, 6)

ws["A15"] = "TIP — rehearse on the fixture before the real data"
ws["A15"].font = Font(name=FONT, sz=11, bold=True)
ws["A16"] = ("In the notebook's PARAMS cell set input_path = "
             '"fixture_physician_scores.csv" and run it end to end first. The '
             "fixture has known answers, so if anything looks odd you know it is "
             "the notebook and not your data. Then swap in the Box path.")
ws.merge_cells("A16:F16")
ws["A16"].font = BODY
ws["A16"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[16].height = 40

ws["A18"] = "CLI vs notebook"
ws["A18"].font = Font(name=FONT, sz=11, bold=True)
ws["A19"] = ("Both call the same function — pipeline.run_pipeline() — so they "
             "cannot give different answers. Use the CLI for scheduled or "
             "repeat runs; use the notebook to investigate and debug. The CLI "
             "exits 3 when blockers are present, so it can gate a job.")
ws.merge_cells("A19:F19")
ws["A19"].font = BODY
ws["A19"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[19].height = 40

ws["A21"] = "Codebase size (live formulas over the File Inventory sheet)"
ws["A21"].font = Font(name=FONT, sz=11, bold=True)
# Row 22 on File Inventory is the TOTAL row, so the data range is 5:21.
metrics_rows = [
    ("Files documented", "=COUNTA('File Inventory'!A5:A21)"),
    ("Total lines of Python", "='File Inventory'!F22"),
    ("Diagnostic checks", "=COUNTA(Checks!A5:A14)"),
    ("Config parameters", "=COUNTA(Parameters!A5:A40)"),
]
for i, (label, formula) in enumerate(metrics_rows, start=22):
    ws.cell(row=i, column=1, value=label).font = BODY
    c = ws.cell(row=i, column=2, value=formula)
    c.font = BODY_B
    c.number_format = "#,##0"

# ===========================================================================
# 2. FILE INVENTORY
# ===========================================================================
ws = sheet("File Inventory", "File inventory",
           "Every file, what it is for, and whether you ever need to open it.")
header(ws, 4, ["File", "Type", "What it does", "Do you edit it?",
               "Depends on", "Lines"])

files = [
    ("PAS_Pipeline.ipynb", "NOTEBOOK — ENTRY POINT",
     "The analysis surface. 38 cells: parameters, schema resolution with fuzzy "
     "column matching, dispersion check, full pipeline run, per-check drill-downs "
     "with charts, reason codes, triage queue, inline PAS cards, threshold sweeps.",
     "Yes — the PARAMS cell. That is the design.",
     "everything below", 0),
    ("pas_diagnostics/config.py", "module",
     "DiagnosticConfig dataclass. Every threshold in the project lives here — 17 "
     "sites in checks.py read from it. Serialises to JSON so a run is reproducible.",
     "No — override from PARAMS instead", "—", 158),
    ("pas_diagnostics/schema.py", "module",
     "ColumnMap and Variable. Maps the 23 scored variables to their raw and score "
     "columns, plus the sentinel registry (-99 etc.), plausible ranges, and the "
     "distributions transcribed from the design workbook.",
     "Only if column names change permanently", "config", 243),
    ("pas_diagnostics/metrics.py", "module",
     "Statistics: Gini, Lorenz, Somers' D, VIF, PSI, weighted quantiles, "
     "band summaries, and monotonicity with a permutation significance test. "
     "pandas + numpy only — no scikit-learn or statsmodels.",
     "No", "—", 291),
    ("pas_diagnostics/checks.py", "module",
     "The ten diagnostic checks. Each returns a table plus severity-tagged "
     "findings. This is where the analytical judgement lives.",
     "Only to add a new check", "metrics, schema, config", 699),
    ("pas_diagnostics/pipeline.py", "module",
     "Orchestration shared by the CLI and notebook: load, resolve schema, fuzzy-"
     "suggest column names, prepare the target, run all checks, return a "
     "PipelineResult.",
     "No", "checks, report, config, schema", 397),
    ("pas_diagnostics/reasons.py", "module",
     "Underwriter tooling: exact additive score attribution, reason codes, "
     "reconciliation, renewal triage queue, composite dispersion.",
     "No", "schema, config", 258),
    ("pas_diagnostics/card.py", "module",
     "Renders the PAS card as standalone HTML — no build step, no CDN. Shows each "
     "driver's raw value beside its score and renders data-quality caveats as "
     "visible warnings.",
     "Only to restyle", "reasons, schema", 242),
    ("pas_diagnostics/report.py", "module",
     "Output writers: markdown findings report, per-check CSVs, optional Excel "
     "workbook.",
     "No", "checks", 143),
    ("pas_diagnostics/cli.py", "module",
     "Command-line entry point. Thin wrapper over pipeline.run_pipeline() so the "
     "CLI and notebook cannot drift apart.",
     "No", "pipeline, config, schema", 138),
    ("pas_diagnostics/__init__.py", "module",
     "Package marker. Exports ColumnMap, Variable, DiagnosticConfig. Without this "
     "file the import fails.",
     "No", "config, schema", 8),
    ("pas_diagnostics/__main__.py", "module",
     "Makes 'python -m pas_diagnostics' work.",
     "No", "cli", 4),
    ("test_edges.py", "TEST — run once",
     "32 adversarial tests. Deliberately tries to break the harness rather than "
     "confirm the happy path. This is your install check.",
     "No", "all modules", 427),
    ("make_fixture.py", "TEST — run once",
     "Generates 60,000 synthetic rows reproducing six pathologies from the design "
     "workbook, so the checks can be verified before touching real data.",
     "No", "—", 310),
    ("build_notebook.py", "generator",
     "Emits PAS_Pipeline.ipynb. The notebook is regenerable rather than a "
     "hand-edited artifact that drifts. Edit here for structural changes.",
     "Only for structural notebook changes", "—", 505),
    ("try_uw.py", "example script",
     "Small worked example of the underwriter tooling outside the notebook. "
     "Reference only.",
     "No", "reasons, card, schema", 35),
    ("README.md", "docs",
     "Install, usage, what each check does, how to read the output, and the open "
     "questions the code surfaced.",
     "—", "—", 0),
]
body(ws, 5, files, [30, 22, 62, 30, 26, 8], wrap_cols=(1, 3, 4, 5),
     fills=lambda i, c, v: (ENTRY_FILL if "ENTRY POINT" in v[1]
                            else (WARN_FILL if v[1].startswith("TEST") else None)))
for r in range(5, 5 + len(files)):
    ws.row_dimensions[r].height = 56
tot = 5 + len(files)
ws.cell(row=tot, column=1, value="TOTAL").font = BODY_B
ws.cell(row=tot, column=6, value=f"=SUM(F5:F{tot - 1})").font = BODY_B
ws.cell(row=tot, column=6).number_format = "#,##0"
ws.cell(row=tot, column=6).border = BOX
ws.cell(row=tot, column=1).border = BOX
ws.cell(row=tot + 2, column=1,
        value="Note: line counts are 0 for the notebook and README because they "
              "are not Python source. Extracted from the code, not typed by hand."
        ).font = SUB

# ===========================================================================
# 3. CHECKS
# ===========================================================================
ws = sheet("Checks", "The ten diagnostic checks",
           "Run in this order by pipeline.run_pipeline(). Schema and data quality "
           "come first because everything after them is meaningless if the columns "
           "are not what the design says.")
header(ws, 4, ["#", "Check", "Question it answers", "Can raise",
               "Notebook cell", "Output table"])
checks_rows = [
    (1, "check_schema", "Are the columns what the design workbook says?",
     "BLOCKER if the composite, target or exposure column is absent",
     "Step 1", "schema"),
    (2, "check_data_quality",
     "Missing, zero, out-of-range values — and critically, sentinels that "
     "received a score. CV_LOSS_FREE_YEARS = -99 means unknown but the "
     "documented bins map it to 10, the worst risk.",
     "BLOCKER when a sentinel scores >= 8", "Step 4", "data_quality"),
    (3, "check_adequacy_is_experience",
     "Is Adequacy measuring the physician's loss experience, or a burn / "
     "on-level factor? 93% of physicians are claim-free, so genuine loss "
     "variables must be near-degenerate at zero.",
     "BLOCKER", "Step 4", "adequacy_reality"),
    (4, "check_circularity",
     "Were the bands cut on the outcome? Specialty and state tiers were cut on "
     "portfolio loss ratio, so lift against loss ratio is circular — and loss "
     "ratio reflects rate adequacy, not only risk.",
     "BLOCKER", "Step 4", "circularity"),
    (5, "check_variable_signal",
     "Does each variable order loss? Band loss ratio vs score, with a "
     "permutation test so noise is not reported as inversion.",
     "BLOCKER if inverted, HIGH if no evidence", "Step 4",
     "variable_signal + bands_*"),
    (6, "check_discrimination",
     "Can the bands separate the book, or is most of it tied? 77% of policies "
     "sit at exactly $1M per-occurrence.",
     "HIGH / MEDIUM", "Step 4", "discrimination"),
    (7, "check_effective_weights",
     "Nominal vs realised weight after weight-zeroing. Hospital Rating is "
     "missing on 59% of records, so its stated weight is fiction for most of "
     "the book.",
     "HIGH", "Step 4", "effective_weights"),
    (8, "check_redundancy",
     "VIF and pairwise rank correlation. Total / indemnity / expense loss cost "
     "and loss ratio share a numerator or denominator by construction.",
     "HIGH / MEDIUM", "Step 4", "vif + score_correlation"),
    (9, "check_composite_lift",
     "Gini in-sample and out-of-time, plus Lorenz curves.",
     "BLOCKER if Gini is negative or under 0.10", "Step 4",
     "lift_summary + lift_lorenz_*"),
    (10, "check_stability",
     "PSI year over year. Band widths are recomputed from the portfolio mean "
     "each run, so part of any shift may be recalibration rather than real "
     "movement.",
     "HIGH above PSI 0.25", "Step 4", "stability"),
]
body(ws, 5, checks_rows, [5, 30, 66, 42, 14, 26], wrap_cols=(3, 4, 6))
for r in range(5, 5 + len(checks_rows)):
    ws.row_dimensions[r].height = 62
add_table(ws, "ChecksTbl", 4, 4 + len(checks_rows), 6)

ws.cell(row=17, column=1,
        value="Severity meaning — BLOCKER: PAS should not influence price or "
              "appear in an underwriter workflow until resolved. HIGH: usable "
              "for internal exploration, not defensible in a filing or model "
              "governance review. MEDIUM: worth fixing. INFO: context.").font = SUB
ws.merge_cells("A17:F17")
ws["A17"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[17].height = 40

# ===========================================================================
# 4. KEY FUNCTIONS
# ===========================================================================
ws = sheet("Key Functions", "Functions you will call directly",
           "Everything else is internal. Extracted from the source, not "
           "transcribed.")
header(ws, 4, ["Module", "Function", "What it does", "Used in"])
funcs = [
    ("pipeline", "load_data(cfg)", "Read the scored CSV or parquet, honouring cfg.nrows.", "Step 1"),
    ("pipeline", "resolve_schema(df, cm)",
     "Which expected columns exist? Returns a frame you can filter on found=False.", "Step 1"),
    ("pipeline", "suggest_column_map(df, cm)",
     "Fuzzy-matches unresolved columns using edit distance and token overlap. "
     "String similarity only — it has no idea what the columns mean, so verify "
     "every suggestion.", "Step 1"),
    ("pipeline", "apply_overrides(cm, overrides)",
     "Patch a ColumnMap from a nested dict. Validates names and raises on typos.", "Step 1"),
    ("pipeline", "run_pipeline(cfg, cm, df)",
     "Runs all ten checks. Returns PipelineResult; writing is separate.", "Step 3"),
    ("PipelineResult", ".table(name) / .bands(var)",
     "Fetch any output table by name.", "Step 4"),
    ("PipelineResult", ".blockers() / .findings_df() / .counts()",
     "Findings as frames, filterable by severity or check.", "Step 3"),
    ("PipelineResult", ".write(outdir)",
     "Write markdown report, CSVs, optional Excel, and run_config.json.", "Step 6"),
    ("reasons", "dispersion(df, cm)",
     "Can the composite populate the high bands? Returns the distribution plus "
     "plain-language notes.", "Step 2"),
    ("reasons", "reconciliation(df, cm)",
     "Proves the attribution is exact. Run before showing anyone reason codes — "
     "if the residual exceeds 0.5 they describe a different model than the one "
     "that produced the score.", "Step 5"),
    ("reasons", "reason_codes(df, cm, top_n)",
     "Top N drivers up and down per physician, with raw values, realised weights, "
     "signed score-point contributions, and data-quality caveats.", "Step 5"),
    ("reasons", "attribute(df, cm)",
     "The raw per-variable contribution matrix behind reason codes.", "Step 5"),
    ("reasons", "triage(df, cm, exposure_col)",
     "Renewal queue ranked on premium at risk, not raw score.", "Step 5"),
    ("card", "render(df, cm, path, limit)",
     "Standalone HTML PAS cards. Displayed inline in the notebook.", "Step 5"),
    ("metrics", "gini / somers_d / monotonicity",
     "Call directly for ad-hoc analysis. monotonicity() includes the permutation "
     "p-value.", "ad hoc"),
    ("config", "DiagnosticConfig(**PARAMS)", "Build a config from the PARAMS dict.", "Setup"),
    ("config", "cfg.replace(**kw)",
     "Copy with overrides — used for the threshold sweeps.", "Sweeps"),
]
body(ws, 5, funcs, [16, 40, 74, 14], wrap_cols=(3,))
for r in range(5, 5 + len(funcs)):
    ws.row_dimensions[r].height = 42
add_table(ws, "FuncsTbl", 4, 4 + len(funcs), 4)

# ===========================================================================
# 5. PARAMETERS
# ===========================================================================
ws = sheet("Parameters", "Config parameters",
           "All set from the notebook's PARAMS cell. Never edit a module to tune "
           "a threshold — it breaks reproducibility from run_config.json.")
header(ws, 4, ["Parameter", "Default", "Controls", "When to change it"])
params = [
    ("input_path", "None", "Path to the scored CSV or parquet.",
     "Always. Pre-filled with your Box path."),
    ("output_dir", "pas_diagnostics_out", "Where reports and tables are written.", "Per run."),
    ("nrows", "None", "Read only the first N rows.",
     "Set to 50,000 while iterating; a 48 MB CSV is slow to parse. None for the "
     "final run."),
    ("target_basis", "loss_ratio",
     "loss_ratio divides by on-levelled premium. loss_cost divides by BCE "
     "exposure and is rate-independent.",
     "Switch to loss_cost to test whether lift survives without the circular "
     "specialty and state tiers. This is the single most informative sweep."),
    ("train_years / test_years", "None",
     "Out-of-time split for the lift check.",
     "Set both, or no split is performed and in-sample lift is not evidence."),
    ("anchor", "4.5",
     "The band the portfolio mean lands in; also the origin for attribution.",
     "Rarely. Note the credibility complement is 5, not 4.5 — that mismatch "
     "is why reconciliation shows a residual."),
    ("credibility_complement", "5.0",
     "Value thin-experience scores are pulled toward.",
     "If the pipeline changes its complement."),
    ("min_exposure_share", "0.005",
     "Bands holding less exposure than this are excluded from monotonicity.",
     "Raise if small bands are producing noisy findings."),
    ("n_permutations", "2000",
     "Draws for the monotonicity significance test.",
     "Raise for stabler p-values; 2000 resolves to about 0.0005."),
    ("p_value_threshold", "0.05",
     "Above this the finding is 'no evidence of ordering', in either direction.",
     "Tighten to 0.01 — about 26 variables are tested, so roughly one false "
     "positive is expected at 0.05."),
    ("inversion_threshold", "-0.30",
     "Spearman below this, and significant, is reported as inverted.", "Rarely."),
    ("weak_signal_threshold", "0.30", "Below this counts as weak ordering.", "Rarely."),
    ("tie_threshold", "0.50",
     "Flag a variable when this share of records get an identical score.",
     "Lower to catch milder concentration."),
    ("min_effective_bands", "3.0", "Flag when fewer bands are effectively in use.", "Rarely."),
    ("circularity_corr_threshold", "0.90",
     "Rank correlation above which an undeclared variable is flagged as "
     "possibly target-derived.", "Rarely."),
    ("vif_severe / vif_moderate", "10.0 / 5.0", "Redundancy thresholds.", "Rarely."),
    ("pairwise_corr_threshold", "0.85", "Flag near-duplicate score pairs.", "Rarely."),
    ("psi_unstable", "0.25", "PSI above which the distribution is called unstable.", "Rarely."),
    ("present_pct_threshold", "0.80",
     "Flag variables available on less than this share of records.",
     "Rarely. Hospital Rating sits at 41%."),
    ("burn_factor_max_zero_share", "0.20",
     "Part of the test for 'this is not loss experience'. A genuine loss "
     "variable is zero for most physicians.",
     "Only with a reason — this drives the most consequential finding."),
    ("burn_factor_min_non_null", "0.98", "Second part of the same test.", "As above."),
    ("burn_factor_max_spread", "0.60", "Third part: relative p5-p95 spread.", "As above."),
    ("expected_zero_claim_share", "0.93",
     "Fallback claim-free share when the claim count column is absent.",
     "If your book's claim-free rate differs."),
    ("spec_drift_tolerance", "0.25",
     "Relative gap from the design workbook's documented distribution that "
     "triggers a drift finding.", "Rarely."),
    ("reason_top_n", "3", "Drivers shown up and down per physician.", "Taste."),
    ("card_limit / card_top_drivers", "25 / 4", "Cards rendered, drivers per card.", "Taste."),
    ("triage_top_n / triage_min_premium", "250 / 0.0", "Size and floor of the review queue.", "Taste."),
    ("drop_nonpositive_exposure", "True", "Filter rows with exposure <= 0.",
     "Set False to inspect raw data when the pipeline reports an empty frame."),
    ("write_csvs / write_markdown / write_excel", "True / True / False",
     "Which outputs .write() produces.", "Enable Excel for circulation."),
    ("verbose", "True", "Progress printing.", "False inside sweeps."),
    ("notes", "{}", "Free-form provenance written into run_metadata.json.",
     "Record analyst and purpose — model governance will ask."),
]
body(ws, 5, params, [34, 22, 60, 58], wrap_cols=(3, 4))
for r in range(5, 5 + len(params)):
    ws.row_dimensions[r].height = 44
add_table(ws, "ParamsTbl", 4, 4 + len(params), 4)

del wb["Sheet"]
out = Path("/mnt/user-data/outputs/PAS_Code_Guide.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"wrote {out}")
print("sheets:", wb.sheetnames)
